"""기획 발표 PPT 생성: python-pptx 필요. 실행 -> docs/tailrisk_기획발표.pptx"""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

DOCS = Path(__file__).parent
FONT = "맑은 고딕"
INK, SLATE, CORAL, TEAL = RGBColor(0x14, 0x21, 0x3D), RGBColor(0x5B, 0x6B, 0x78), RGBColor(0xD6, 0x45, 0x50), RGBColor(0x2A, 0x9D, 0x8F)
TXT, PALE, WHITE, LIGHT = RGBColor(0x1B, 0x26, 0x30), RGBColor(0xEE, 0xF1, 0xF5), RGBColor(255, 255, 255), RGBColor(0xB8, 0xC4, 0xD6)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, s, size=16, bold=False, color=TXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(s if isinstance(s, list) else [s]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    return tb


def box(slide, x, y, w, h, fill=PALE, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.06
    return s


def title(slide, s, sub=None):
    text(slide, 0.7, 0.5, 11.9, 0.8, s, 30, True, INK)
    if sub:
        text(slide, 0.7, 1.25, 11.9, 0.5, sub, 16, False, SLATE)


def page(n, notes):
    sl = prs.slides.add_slide(BLANK)
    bg(sl, WHITE)
    text(sl, 12.0, 7.0, 0.7, 0.3, str(n), 11, False, SLATE, PP_ALIGN.RIGHT)
    sl.notes_slide.notes_text_frame.text = notes
    return sl


def badge(slide, x, y, n, fill=INK, d=0.6):
    box(slide, x, y, d, d, fill, MSO_SHAPE.OVAL)
    text(slide, x, y, d, d, str(n), 18, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


# 1. 표지
s = prs.slides.add_slide(BLANK)
bg(s, INK)
text(s, 0.9, 2.1, 11.5, 1.2, "tailrisk", 60, True, WHITE)
text(s, 0.9, 3.4, 11.5, 1.0, "변동성 급등을 앞서 감지하고, 비중을 조절하는 리스크 타겟팅 실험", 26, False, LIGHT)
text(s, 0.9, 4.7, 11.5, 0.5, "LightGBM은 HAR-RV와 GARCH를 이길 수 있을까", 18, False, CORAL)
text(s, 0.9, 6.5, 11.5, 0.4, "기획 발표  |  2026.10  |  기한 10.19", 14, False, LIGHT)
s.notes_slide.notes_text_frame.text = "무엇을 묻고, 어떻게 검증하고, 어떤 일정으로 진행하는지를 설명합니다. 결과 발표가 아니라 기획 발표입니다."

# 2. 왜
s = page(2, "변동성 군집은 실증적으로 잘 알려진 성질입니다. 불안한 날 뒤에 불안한 날이 온다는 뜻입니다.")
title(s, "폭락은 한꺼번에 오고, 고정 비중은 그때 가장 아프다", "변동성은 군집한다: 흔들리는 시기가 이어진다")
box(s, 0.7, 2.1, 5.8, 4.4)
text(s, 1.0, 2.35, 5.2, 0.5, "고정 비중", 22, True, SLATE)
text(s, 1.0, 3.1, 5.2, 3.2, ["주식 60, 채권 40처럼 비중을 정해 두고 유지", "시장이 조용하든 요동치든 같은 크기로 위험을 진다",
                             "급락기에 위험이 가장 커져 최대낙폭이 깊어진다"], 17)
box(s, 6.83, 2.1, 5.8, 4.4, INK)
text(s, 7.13, 2.35, 5.2, 0.5, "리스크 타겟팅", 22, True, CORAL)
text(s, 7.13, 3.1, 5.2, 3.2, ["앞으로의 변동성을 예측해 위험 크기를 일정하게 맞춤",
                              "변동성이 클 것 같으면 비중을 줄이고, 작을 것 같으면 늘림",
                              "핵심은 예측이 얼마나 정확한가"], 17, False, WHITE)

# 3. 질문
s = page(3, "질문 하나를 둘로 나눠 검증합니다. 1번이 안 되면 2번도 의미가 약해집니다. 그래서 중단 규칙을 1번 기준으로 둡니다.")
title(s, "질문은 두 개이고, 하나로 이어진다")
for i, (n, h, d, m) in enumerate([
        (1, "예측", "LightGBM이 HAR-RV와 GARCH보다 앞으로 20영업일 실현변동성을 더 잘 맞히는가", "지표: QLIKE (보조 MSE)"),
        (2, "배분", "그 예측으로 목표 변동성(연 10%) 배분을 하면 고정 비중보다 최대낙폭이 줄어드는가", "지표: 최대낙폭 (보조 샤프, 회전율)")]):
    x = 0.7 + i * 6.1
    box(s, x, 1.9, 5.8, 4.2)
    badge(s, x + 0.35, 2.2, n, CORAL if i else INK)
    text(s, x + 1.2, 2.3, 4.2, 0.5, h, 24, True, INK)
    text(s, x + 0.35, 3.2, 5.1, 1.8, d, 17)
    text(s, x + 0.35, 5.2, 5.1, 0.5, m, 14, False, SLATE)
text(s, 0.7, 6.4, 11.9, 0.4, "실현변동성: 최근 일별 수익률의 흔들림 크기를 연율로 환산한 값", 13, False, SLATE)

# 4. 접근 흐름
s = page(4, "배분 규칙은 단순합니다. 목표 위험을 예측 변동성으로 나눈 값이 비중입니다. 예를 들어 목표 10%인데 예측이 20%면 비중은 절반입니다.")
title(s, "데이터에서 비중까지 한 줄로 이어진다", "비중 = 목표 변동성 10% ÷ 예측 변동성")
flow = ["수집\n야후·FRED·ECOS", "피처·라벨\n시점별 값만", "모델 3종\nHAR·GARCH·LGBM", "예측\n20일 변동성", "비중\n목표 10% 맞춤", "비교\n고정 비중 대비"]
for i, f in enumerate(flow):
    x = 0.7 + i * 2.02
    box(s, x, 2.5, 1.8, 2.0, INK if i in (2, 4) else PALE)
    h, d = f.split("\n")
    c = WHITE if i in (2, 4) else INK
    text(s, x + 0.15, 2.8, 1.5, 0.5, h, 18, True, c, PP_ALIGN.CENTER)
    text(s, x + 0.15, 3.5, 1.5, 0.9, d, 12, False, LIGHT if i in (2, 4) else SLATE, PP_ALIGN.CENTER)
    if i < 5:
        text(s, x + 1.8, 3.2, 0.22, 0.5, "›", 22, True, CORAL, PP_ALIGN.CENTER)
text(s, 0.7, 5.2, 11.9, 1.2, ["자산마다 독립된 모델과 독립된 타겟팅을 씁니다.", "비용은 왕복 10bp(0.1%)로 가정해 배분 단계에서 반영합니다."], 16)

# 5. 모델
s = page(5, "HAR-RV는 일·주·월 변동성을 섞어 쓰는 단순 회귀입니다. GARCH는 어제 충격이 오늘 변동성을 키우는 구조를 모델링합니다. 둘 다 오래 검증된 기준선입니다.")
title(s, "단순한 기준선 둘에 머신러닝 하나를 붙인다")
ms = [("HAR-RV", "기준선", "최근 하루·1주·1달 변동성을 섞어 앞날을 예측하는 단순 회귀. 금융에서 가장 강한 기본기", SLATE),
      ("GARCH(1,1)", "기준선", "어제의 충격과 어제의 변동성으로 오늘의 변동성을 설명하는 통계 모델", SLATE),
      ("LightGBM", "도전자", "VIX, 금리, 신용스프레드 등 많은 신호를 한꺼번에 쓰는 트리 기반 머신러닝", CORAL)]
for i, (n, tag, d, col) in enumerate(ms):
    x = 0.7 + i * 4.1
    box(s, x, 1.9, 3.8, 4.3)
    text(s, x + 0.3, 2.15, 3.2, 0.4, tag, 14, True, col)
    text(s, x + 0.3, 2.65, 3.2, 0.6, n, 24, True, INK)
    text(s, x + 0.3, 3.55, 3.2, 2.5, d, 15)
text(s, 0.7, 6.4, 11.9, 0.4, "LSTM은 하지 않습니다. QLIKE: 변동성 예측 오차를 재는 지표로 낮을수록 좋습니다.", 13, False, SLATE)

# 6. 범위
s = page(6, "범위를 줄인 것이 이번 기획의 핵심 결정입니다. 기한 안에 질문 하나를 끝까지 답하는 쪽을 택했습니다.")
title(s, "범위를 좁혀서 질문 하나에 끝까지 답한다")
box(s, 0.7, 1.9, 5.8, 4.5)
text(s, 1.0, 2.1, 5.2, 0.5, "한다", 22, True, TEAL)
text(s, 1.0, 2.9, 5.2, 3.4, ["자산 5개: S&P500, KOSPI, 미국 장기채(TLT), 금(GLD), 달러원", "회귀 하나: 20영업일 실현변동성",
                              "모델 3개, 지표 2개", "MLflow로 폴드별 실험 기록"], 16)
box(s, 6.83, 1.9, 5.8, 4.5)
text(s, 7.13, 2.1, 5.2, 0.5, "하지 않는다", 22, True, CORAL)
text(s, 7.13, 2.9, 5.2, 3.4, ["LSTM 등 딥러닝", "VaR 백테스트 (Kupiec 등)", "레짐 진입 분류: 회귀 결과를 임계값으로 자른 부록으로만",
                              "여러 문제를 동시에 푸는 것"], 16)

# 7. 데이터
s = page(7, "일정과 무관하게 필요한 데이터를 전부 모으기로 결정했습니다. 하이일드 스프레드는 FRED 정책상 최근 3년만 있어 무디스 스프레드와 HYG/IEF 비율로 대체합니다.")
title(s, "2000년부터 약 50개 열을 하나의 일별 표로 모았다", "달력은 S&P500 거래일, 2000-01-03부터. 수집 완료")
cd = CategoryChartData()
cats = ["S&P500·KOSPI·VIX·DXY·금리", "MOVE·TLT·IEF·LQD", "달러원 (야후)", "금 (GLD)", "VIX 3개월", "하이일드 (HYG)", "VIX 9일"]
cd.categories = cats
cd.add_series("시작 연도", (2000, 2002, 2003, 2004, 2006, 2007, 2011))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.7), Inches(2.0), Inches(7.4), Inches(4.7), cd)
ch = gf.chart
ch.has_legend, ch.font.size, ch.font.name = False, Pt(13), FONT
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.minimum_scale = 1995
ch.category_axis.reverse_order = True
pl = ch.plots[0]
pl.has_data_labels = True
pl.data_labels.number_format, pl.data_labels.number_format_is_linked = "0", False
pl.data_labels.font.size = Pt(13)
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.series[0].format.fill.solid()
pl.series[0].format.fill.fore_color.rgb = TEAL
box(s, 8.5, 2.0, 4.1, 4.7)
text(s, 8.8, 2.25, 3.5, 4.3, ["출처 세 곳", "야후: 자산 5종, VIX 기간구조, MOVE, 달러지수, 신용 ETF", "FRED: 금리, 무디스 신용스프레드 (API 키 경로)",
                               "ECOS: 국고채, CD, 기준금리, 외국인 순매수, 원달러 15:30"], 14)

# 8. 평가 원칙
s = page(8, "미래 원자료를 잘라내도 과거 예측이 변하지 않는지 테스트로 확인합니다. 이 테스트가 누수 방지의 증거입니다.")
title(s, "미래 정보를 쓰지 않는 것이 첫째 원칙이다")
rules = [("t 시점엔 t까지만", "FRED는 1영업일 지연을 일괄 적용. 한국 지표는 당일 값"),
         ("학습 표본은 끝난 것만", "목표 구간이 t 이전에 끝난 시점(s ≤ t-20)만 학습에 사용"),
         ("walk-forward", "2010-01부터, 확장 윈도우, 5영업일 간격 예측, 약 6개월마다 재학습"),
         ("테스트로 확인", "미래 원자료를 잘라내도 과거 예측이 같은지 자동 검증")]
for i, (h, d) in enumerate(rules):
    x, y = 0.7 + (i % 2) * 6.1, 1.8 + (i // 2) * 2.5
    box(s, x, y, 5.8, 2.2)
    text(s, x + 0.3, y + 0.25, 5.2, 0.5, h, 20, True, INK)
    text(s, x + 0.3, y + 0.95, 5.2, 1.2, d, 15)

# 9. 피처
s = page(9, "피처를 두 단계로 나눈 이유는 늦게 시작하는 신호가 정말 도움이 되는지 따로 재기 위해서입니다.")
title(s, "피처는 두 단계로 넣어 기여를 따로 잰다")
box(s, 0.7, 1.9, 5.8, 4.4)
badge(s, 1.0, 2.15, 1)
text(s, 1.85, 2.25, 4.4, 0.5, "1차: 2000년부터 있는 것", 20, True, INK)
text(s, 1.0, 3.1, 5.2, 3.1, ["자산별 실현변동성·수익률", "VIX, 달러지수", "미국·한국 금리, 무디스 스프레드", "달러원 15:30 종가"], 16)
box(s, 6.83, 1.9, 5.8, 4.4, INK)
badge(s, 7.13, 2.15, 2, CORAL)
text(s, 7.98, 2.25, 4.4, 0.5, "2차: 늦게 시작하는 것", 20, True, WHITE)
text(s, 7.13, 3.1, 5.2, 3.1, ["HYG/IEF 신용 비율 (2007~)", "VIX 3개월, MOVE", "외국인 순매수 (2003~)", "VIX 9일은 본 실험 제외, 부록 비교만"], 16, False, WHITE)
text(s, 0.7, 6.5, 11.9, 0.4, "모든 피처는 t 시점에 알 수 있는 값만 사용합니다.", 13, False, SLATE)

# 10. MLOps
s = page(10, "MLOps는 핵심 결과가 나온 뒤에 감쌉니다. Windows라 Airflow는 Docker가 필요하고, 부담이 크면 Prefect로 갑니다.")
title(s, "실험은 기록하고, 마지막에 일별 파이프라인으로 감싼다")
box(s, 0.7, 1.9, 5.8, 4.4)
text(s, 1.0, 2.1, 5.2, 0.5, "MLflow", 22, True, INK)
text(s, 1.0, 2.9, 5.2, 3.2, ["폴드별 run을 기록", "모델, 피처 묶음, QLIKE를 한 곳에서 비교", "나중에 같은 결과를 다시 만들 수 있게 남김"], 16)
box(s, 6.83, 1.9, 5.8, 4.4, INK)
text(s, 7.13, 2.1, 5.2, 0.5, "일별 DAG (Airflow 또는 Prefect)", 20, True, WHITE)
for i, st in enumerate(["수집", "피처", "예측", "비중"]):
    x = 7.13 + i * 1.35
    box(s, x, 3.2, 1.15, 0.9, WHITE)
    text(s, x, 3.2, 1.15, 0.9, st, 16, True, INK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 7.13, 4.5, 5.2, 1.6, "핵심 결과 뒤에 감싸기만 합니다.", 15, False, LIGHT)

# 11. 일정
s = page(11, "10월 12일까지 HAR-RV 대 LightGBM 비교가 안 나오면 배분과 DAG를 줄이고 비교만 끝내는 중단 규칙을 뒀습니다.")
title(s, "10월 19일까지, 비교를 먼저 끝낸다", "10월 12일까지 모델 비교가 안 나오면 배분과 DAG를 줄이고 비교만 마무리")
sched = [("10/08", "뼈대, 데이터 수집, 패널 규칙 테스트", "완료"), ("10/09~10", "피처·라벨, walk-forward 틀, HAR-RV 기준선", "진행"),
         ("10/11~12", "GARCH, LightGBM, MLflow 기록 시작", "다음"), ("10/13~14", "변동성 타겟 배분, 고정 비중 비교", ""),
         ("10/15~16", "일별 DAG로 감싸기", ""), ("10/17~18", "문서, 그림, 재현 확인", ""), ("10/19", "여유분", "")]
for i, (d, t, st) in enumerate(sched):
    y = 2.0 + i * 0.67
    box(s, 0.7, y, 11.9, 0.57, INK if st == "다음" else PALE)
    c = WHITE if st == "다음" else INK
    text(s, 1.0, y, 1.8, 0.57, d, 15, True, CORAL if st == "다음" else INK, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.9, y, 7.6, 0.57, t, 15, False, c, anchor=MSO_ANCHOR.MIDDLE)
    if st:
        text(s, 10.8, y, 1.5, 0.57, st, 14, True, TEAL if st == "완료" else CORAL, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)

# 12. 한계 + 마무리
s = page(12, "LightGBM이 지는 결과도 의미 있는 결과로 기록합니다. 질문에 정직하게 답하는 것이 목표입니다.")
bg(s, INK)
text(s, 0.7, 0.6, 11.9, 0.9, "미리 밝혀 두는 한계", 30, True, WHITE)
text(s, 0.7, 1.9, 11.9, 3.2, ["야후 데이터는 과거 수정이 있을 수 있어 완전 재현을 보장하지 못합니다.",
                               "FRED 지연을 1영업일로 일괄 적용했지만, 실제 발표 시각은 지표마다 다릅니다.",
                               "거래비용은 배분 단계에서 왕복 10bp 가정으로만 반영합니다.",
                               "하이일드 스프레드 원계열은 최근 3년뿐이라 대용 지표를 씁니다."], 18, False, WHITE)
text(s, 0.7, 5.4, 11.9, 1.0, "LightGBM이 이기든 지든, 기준선과 같은 조건에서 비교한 결과를 그대로 기록합니다.", 20, True, CORAL)

out = DOCS / "tailrisk_기획발표.pptx"
prs.save(out)
print(out)
